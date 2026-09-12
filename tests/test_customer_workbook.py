"""The customer-facing weekly Excel workbook.

Two things are being protected here. The first is that the workbook is usable:
three tabs, a real table with filters, frozen identity columns, a status
dropdown and working links. The second matters more — that nothing internal or
unsupported reaches a customer. No local file paths, no confidence figures, no
taxonomy keys, no invented product detail, and no wording that turns a
commercial signal into a claim that somebody asked to be sold to.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.deliver.customer_narrative import (
    build_report,
    build_row,
    priority_for,
    summarise_product,
    why_now,
)
from src.deliver.customer_workbook import (
    BRIEF_SHEET,
    HOW_TO_SHEET,
    OPPORTUNITIES_SHEET,
    write_workbook,
)
from src.sales.matching import MatchedLead, select_for_prospect
from src.settings import load_config
from tests.conftest import make_lead, make_prospect

PUBLISHED = date(2025, 12, 12)

# Wording that would mean we had leaked internals or overstated the product.
FORBIDDEN_FRAGMENTS = [
    "file://",
    "/home/",
    "C:\\",
    "confidence",
    "fit_score",
    "dedupe",
    "nice class",
    "sic code",
    "launchtrace_score",
    "cereal_bars",
    "flexible_packaging",
]


@pytest.fixture
def config():  # type: ignore[no-untyped-def]
    return load_config("customer_report.json")


def _match(lead=None, fit: float = 18.0) -> MatchedLead:  # type: ignore[no-untyped-def]
    return MatchedLead(lead=lead or make_lead(), fit_score=fit, reasons=[], supplier_line="")


def _report(config, leads=None):  # type: ignore[no-untyped-def]
    prospect = make_prospect(company_name="Digimock")
    result = select_for_prospect(
        prospect,
        leads or [make_lead()],
        source="test",
        count=50,
        reference_date=PUBLISHED,
    )
    report = build_report(result, config=config, highlight_count=4)
    report.journal_number = "2025-050"
    return report


def _all_text(path: Path) -> str:
    workbook = load_workbook(path, rich_text=True)
    chunks: list[str] = []
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    chunks.append(str(cell.value))
                if cell.hyperlink is not None:
                    chunks.append(str(cell.hyperlink.target))
    return "\n".join(chunks)


class TestCustomerLanguage:
    def test_a_product_line_only_names_formats_the_filing_actually_lists(self, config):
        lead = make_lead(
            product_category="coffee_tea",
            product_category_label="Coffee, tea and hot drinks",
            goods_summary="Coffee; Ground coffee; Coffee beans; Drip bag coffee.",
        )
        head, line = summarise_product(_match(lead), config)
        assert head == "Coffee"
        assert "ground" in line and "drip bags" in line
        # Nothing about pods: the filing never mentions them.
        assert "pods" not in line and "capsules" not in line

    def test_a_goods_list_is_compressed_not_reproduced(self, config):
        lead = make_lead(
            product_category="snacks",
            product_category_label="Snacks",
            goods_summary="; ".join(["Roasted nuts", "Salted nuts", "Flavoured nuts"] * 60),
        )
        _, line = summarise_product(_match(lead), config)
        assert len(line) < 120, "a long goods list must not become a huge row"

    def test_an_uncategorised_filing_falls_back_rather_than_guessing(self, config):
        lead = make_lead(product_category="", product_category_label="Not categorised")
        head, line = summarise_product(_match(lead), config)
        assert head == line == "Not categorised"

    def test_why_now_states_the_company_age_the_score_already_recorded(self, config):
        lead = make_lead(
            company_incorporation_date=date(2025, 3, 1),
            filing_date=date(2025, 9, 15),
            reasons=["First trade mark we have seen from this applicant"],
        )
        text = why_now(_match(lead))
        assert "6 months before the filing" in text
        assert "First trade mark" in text

    def test_an_unconfirmed_website_is_said_plainly_not_invented(self, config):
        lead = make_lead(website="")
        row = build_row(_match(lead), config, "flexible_packaging", {})
        assert row.website == "Not confirmed"
        assert row.website_url == ""
        assert "No website we could confirm" in row.why_now

    def test_priority_is_customer_facing_and_never_the_raw_band(self, config):
        assert priority_for(_match(make_lead(band="HIGH")), config) == "TOP MATCH"
        assert priority_for(_match(make_lead(band="MEDIUM"), fit=15.0), config) == "STRONG"
        assert priority_for(_match(make_lead(band="MEDIUM"), fit=4.0), config) == "RELEVANT"

    def test_the_angle_is_framed_as_a_suggestion(self, config):
        disclaimer = config["copy"]["angle_disclaimer"].lower()
        assert "not requests" in disclaimer
        how_to = " ".join(block["body"] for block in config["copy"]["how_to_use"]).lower()
        assert "not as an inbound enquiry" in how_to
        assert "nothing here is a confirmed buying request" in how_to


class TestReportAssembly:
    def test_every_qualifying_opportunity_reaches_the_sheet(self, config):
        leads = [
            make_lead(),
            make_lead(
                brand_name="SECOND",
                trademark_number="UK00003900002",
                company_number="14000002",
                company_name="SECOND FOODS LTD",
                product_category="coffee_tea",
                product_category_label="Coffee, tea and hot drinks",
                goods_summary="Coffee; Ground coffee.",
            ),
        ]
        report = _report(config, leads)
        assert report.total == 2
        assert {row.brand for row in report.rows} == {"CRUMBLEDGE", "SECOND"}

    def test_the_brief_features_the_matcher_s_own_first_picks(self, config):
        report = _report(config)
        assert [row.brand for row in report.highlights] == [row.brand for row in report.rows][:4]

    def test_a_flat_week_is_described_as_flat(self, config):
        report = _report(config)
        assert "stands out" in report.signal


class TestWorkbookFile:
    @pytest.fixture
    def path(self, config, tmp_path: Path) -> Path:
        return write_workbook(_report(config), tmp_path / "workbook.xlsx", config=config)

    def test_it_has_the_three_tabs(self, path):
        assert load_workbook(path).sheetnames == [
            BRIEF_SHEET,
            OPPORTUNITIES_SHEET,
            HOW_TO_SHEET,
        ]

    def test_the_working_sheet_is_a_filterable_table_with_frozen_identity_columns(self, path):
        sheet = load_workbook(path)[OPPORTUNITIES_SHEET]
        table = sheet.tables["LaunchTraceOpportunities"]
        assert table.autoFilter is not None
        assert sheet.freeze_panes == "D5"
        assert sheet.sheet_view.showGridLines is False

    def test_sales_status_offers_a_dropdown(self, path, config):
        sheet = load_workbook(path)[OPPORTUNITIES_SHEET]
        validations = sheet.data_validations.dataValidation
        assert len(validations) == 1
        for option in config["sales_status_options"]:
            assert option in validations[0].formula1

    def test_the_three_customer_columns_are_left_empty(self, path):
        sheet = load_workbook(path)[OPPORTUNITIES_SHEET]
        headers = [cell.value for cell in sheet[4]]
        assert headers[-3:] == ["Owner", "Sales Status", "Sales Notes"]
        for row in sheet.iter_rows(min_row=5, min_col=len(headers) - 2):
            assert all(cell.value is None for cell in row)

    def test_no_internal_columns_are_exposed(self, path):
        headers = {str(cell.value) for cell in load_workbook(path)[OPPORTUNITIES_SHEET][4]}
        assert not headers & {"LaunchTrace Score", "Fit", "Band", "Confidence", "Nice Classes"}

    def test_every_link_is_a_public_https_url(self, path):
        workbook = load_workbook(path)
        links = [
            cell.hyperlink.target
            for sheet in workbook
            for row in sheet.iter_rows()
            for cell in row
            if cell.hyperlink
        ]
        assert links, "the sheet should carry clickable sources"
        assert all(link.startswith("https://") for link in links)

    def test_the_workbook_carries_no_formulas(self, path):
        workbook = load_workbook(path)
        formulas = [
            cell.coordinate
            for sheet in workbook
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert formulas == []

    def test_nothing_internal_or_local_leaks_into_the_file(self, path):
        text = _all_text(path).lower()
        for fragment in FORBIDDEN_FRAGMENTS:
            assert fragment.lower() not in text, f"{fragment} reached a customer-facing sheet"

    def test_it_says_it_is_a_sample(self, path):
        assert "Sample — not a live subscription" in _all_text(path)

    def test_no_row_becomes_enormous(self, path):
        sheet = load_workbook(path)[OPPORTUNITIES_SHEET]
        heights = [
            dimension.height
            for dimension in sheet.row_dimensions.values()
            if dimension.height is not None
        ]
        assert heights and max(heights) <= 150
